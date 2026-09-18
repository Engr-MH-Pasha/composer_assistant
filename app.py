import io
import base64
import streamlit as st
import pandas as pd
from pypdf import PdfReader
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import arabic_reshaper
from bidi.algorithm import get_display
from groq import Groq

import auth

# Page Configuration
st.set_page_config(page_title="Composer Assistant", layout="wide", page_icon="✍️")

# Enforce Authentication
if not auth.check_password():
    st.title("✍️ Composer Assistant (اردو اور انگریزی)")
    st.info("👈 Please enter your User ID and Password in the sidebar to unlock the workspace.")
    st.stop()

# Show logged-in bar
auth.render_logout()

# Initialize Groq Client
GROQ_API_KEY = st.secrets.get("GROQ_API_KEY", "")
if not GROQ_API_KEY:
    st.error("Groq API key not configured in secrets.toml / Streamlit Secrets.")
    st.stop()

client = Groq(api_key=GROQ_API_KEY)

# Helper: Extract raw text from files
def extract_text_from_file(uploaded_file):
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        return "\n".join([page.extract_text() or "" for page in reader.pages])
    elif name.endswith(".docx"):
        doc = Document(uploaded_file)
        return "\n".join([p.text for p in doc.paragraphs])
    elif name.endswith((".xlsx", ".xls")):
        df = pd.read_excel(uploaded_file)
        return df.to_string(index=False)
    return ""

# Helper: Groq LLM Processing Engine
def run_composition(raw_content: str, mode: str) -> str:
    prompts = {
        "Original": (
            "You are an exact proofreader and grammatical corrector for Urdu and English. "
            "Fix ONLY spelling mistakes, typographical errors, and minor sentence grammatical structures. "
            "Keep the original wording, tone, style, and structure completely intact. "
            "Do NOT add headings, bullet points, or new sentences unless present in the raw input. "
            "Output the clean text directly without conversational disclaimers."
        ),
        "Default": (
            "You are a professional text composer for Urdu and English. "
            "Standardize punctuation, fix syntax, clean paragraph flow, and compose the text neatly. "
            "Maintain the exact intent and information provided without dramatic expansions. "
            "Output the clean composed text directly without any conversational disclaimers."
        ),
        "Improver": (
            "You are an executive editor and formatting specialist for Urdu and English. "
            "Transform this content into an executive, highly polished, standard layout. "
            "Use clear headings, structured paragraphs, bullet points, and appropriate formal vocabulary. "
            "Ensure the output reflects official enterprise documentation standards. "
            "Output the final polished document directly without conversational disclaimers."
        )
    }

    system_prompt = prompts.get(mode, prompts["Default"])
    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": raw_content}
        ],
        temperature=0.2
    )
    return response.choices[0].message.content

# Helper: Export Functions
def export_to_docx(text: str) -> io.BytesIO:
    doc = Document()
    for line in text.split("\n"):
        p = doc.add_paragraph(line)
        # Rudimentary check for Urdu RTL script
        if any("\u0600" <= char <= "\u06FF" for char in line):
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

def export_to_pdf(text: str) -> io.BytesIO:
    bio = io.BytesIO()
    pdf = SimpleDocTemplate(bio, pagesize=letter)
    styles = getSampleStyleSheet()
    normal_style = styles['Normal']
    story = []

    for line in text.split("\n"):
        if not line.strip():
            story.append(Spacer(1, 10))
            continue
        # Reshape Urdu strings to prevent disconnected/reversed characters
        if any("\u0600" <= char <= "\u06FF" for char in line):
            reshaped_text = arabic_reshaper.reshape(line)
            bidi_text = get_display(reshaped_text)
            story.append(Paragraph(bidi_text, normal_style))
        else:
            story.append(Paragraph(line, normal_style))
        story.append(Spacer(1, 6))

    pdf.build(story)
    bio.seek(0)
    return bio

def export_to_excel(text: str) -> io.BytesIO:
    bio = io.BytesIO()
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    df = pd.DataFrame({"Composed Content": lines})
    with pd.ExcelWriter(bio, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name="Composed")
    bio.seek(0)
    return bio

# ==========================================
# Main User Interface
# ==========================================
st.title("✍️ Composer Assistant")
st.caption("Bilingual Urdu & English Document Processing and Composition Suite")

# Mode Selection
st.subheader("1. Composition Setting")
mode = st.radio(
    "Choose Processing Mode:",
    options=["Original", "Default", "Improver"],
    index=1,
    horizontal=True,
    help=(
        "Original: Corrects only spelling/syntax while leaving structure 100% untouched.\n"
        "Default: Composes clean, standard text.\n"
        "Improver: Fully restructures into executive enterprise layout."
    )
)

st.markdown("---")
st.subheader("2. Select Input Method")

tab1, tab2, tab3, tab4 = st.tabs([
    "📁 Document Upload", 
    "🎙️ Voice Recording", 
    "🖼️ Image OCR", 
    "📝 Direct Writing"
])

raw_extracted_text = ""

with tab1:
    doc_files = st.file_uploader(
        "Upload PDF, Word (.docx), or Excel (.xlsx)", 
        type=["pdf", "docx", "xlsx"], 
        accept_multiple_files=True  # 👈 اس سے متعدد فائلیں اپلوڈ ہوں گی
    )
    if doc_files:
        combined_texts = []
        for file in doc_files:
            file_text = extract_text_from_file(file)
            if file_text.strip():
                combined_texts.append(f"--- Document: {file.name} ---\n{file_text}")
                st.success(f"Loaded: {file.name}")
        raw_extracted_text = "\n\n".join(combined_texts)
with tab2:
    st.write("Record spoken Urdu or English:")
    audio_data = st.audio_input("Record Voice")
    if audio_data:
        st.info("Transcribing audio via Groq Whisper...")
        transcription = client.audio.transcriptions.create(
            file=("audio.wav", audio_data.read()),
            model="whisper-large-v3",
            response_format="text"
        )
        raw_extracted_text = str(transcription)
        st.success("Audio transcribed successfully!")

with tab3:
    image_files = st.file_uploader(
        "Upload document snapshots or handwritten notes (JPG / PNG)", 
        type=["jpg", "jpeg", "png"], 
        accept_multiple_files=True  # 👈 اس سے متعدد تصاویر اپلوڈ ہوں گی
    )
    if image_files:
        st.write(f"کل تصاویر منتخب ہوئیں: {len(image_files)}")
        if st.button("Extract Text from All Images"):
            all_image_texts = []
            progress_bar = st.progress(0)
            
            for idx, img in enumerate(image_files):
                base64_image = base64.b64encode(img.getvalue()).decode('utf-8')
                mime = img.type
                
                vision_response = client.chat.completions.create(
                    model="llama-3.2-11b-vision-preview",
                    messages=[{
                        "role": "user",
                        "content": [
                            {"type": "text", "text": "Extract and transcribe all written Urdu and English text from this image exactly as written. Provide only the text."},
                            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{base64_image}"}}
                        ]
                    }],
                    temperature=0.1
                )
                img_text = vision_response.choices[0].message.content
                all_image_texts.append(f"--- Image {idx+1}: {img.name} ---\n{img_text}")
                progress_bar.progress((idx + 1) / len(image_files))
            
            raw_extracted_text = "\n\n".join(all_image_texts)
            st.session_state["extracted_img_text"] = raw_extracted_text
            st.success("تمام تصاویر سے ٹیکسٹ نکال لیا گیا!")

    # اگر بٹن دبنے کے بعد پیج ریفریش ہو تو ٹیکسٹ برقرار رہے
    if "extracted_img_text" in st.session_state and not raw_extracted_text:
        raw_extracted_text = st.session_state["extracted_img_text"]
with tab4:
    direct_text = st.text_area("Write or paste raw draft here:", height=180)
    if direct_text.strip():
        raw_extracted_text = direct_text

# ==========================================
# Execution & Output Generation
# ==========================================
if raw_extracted_text:
    st.markdown("---")
    with st.expander("Show Extracted Raw Input"):
        st.text(raw_extracted_text)

    if st.button(f"Generate Composition ({mode} Mode)", type="primary"):
        with st.spinner("Processing through Groq LLM..."):
            st.session_state["composed_output"] = run_composition(raw_extracted_text, mode)

if "composed_output" in st.session_state:
    st.subheader("3. Composed Result")
    composed_text = st.session_state["composed_output"]
    st.text_area("Composed Text Output", composed_text, height=280)

    # Export Area
    st.subheader("4. Download Formats")
    col1, col2, col3 = st.columns(3)

    with col1:
        docx_bytes = export_to_docx(composed_text)
        st.download_button(
            label="📄 Download Word (.docx)",
            data=docx_bytes,
            file_name="Composed_Document.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )

    with col2:
        pdf_bytes = export_to_pdf(composed_text)
        st.download_button(
            label="📑 Download PDF (.pdf)",
            data=pdf_bytes,
            file_name="Composed_Document.pdf",
            mime="application/pdf",
            use_container_width=True
        )

    with col3:
        excel_bytes = export_to_excel(composed_text)
        st.download_button(
            label="📊 Download Excel (.xlsx)",
            data=excel_bytes,
            file_name="Composed_Document.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
